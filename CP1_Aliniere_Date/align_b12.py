"""CP1 U12: aliniere RNMCA B-12 cu Open-Meteo pe zi locala Europe/Bucharest.
Rulare: python align_b12.py --rnmca RNMCA_B12_concentratii.xlsx --meteo 'Meteo.csv(1).xlsx'
Necesita: openpyxl. Raportul .xls original trebuie exportat in .xlsx fara modificari.
"""
import argparse,csv,json,collections,datetime as dt,statistics
from zoneinfo import ZoneInfo
from pathlib import Path
from openpyxl import load_workbook

def number(v):
    if v is None or str(v).strip() in ('','-','--'):return None
    try:return float(str(v).replace(',','.'))
    except (ValueError,TypeError):return None

def read_rnmca(path):
    sh=load_workbook(path,read_only=True,data_only=True).active
    d=collections.defaultdict(list); seen=0; invalid=0; duplicate=0; keys=set(); first=None;last=None
    for row in sh.iter_rows(min_row=12,values_only=True):
        t=row[0]
        if not isinstance(t,str):continue
        try:
            day, hour=t.rsplit(' ',1); day=dt.date.fromisoformat(day.replace('/','-'));hour=int(hour)
            if not 1<=hour<=24:continue
        except (ValueError,TypeError):continue
        key=(day,hour)
        if key in keys:duplicate+=1;continue
        keys.add(key);seen+=1;first=min(first,day) if first else day;last=max(last,day) if last else day
        val=number(row[4] if len(row)>4 else None)
        if val is None:invalid+=1
        else:d[day].append(val)
    return d,{'total_hourly_records':seen,'pm25_missing_or_invalid':invalid,'duplicate_hourly_records':duplicate,'first_date':str(first),'last_date':str(last)}

def read_meteo(path):
    sh=load_workbook(path,read_only=True,data_only=True).active
    d=collections.defaultdict(lambda:{'temp':[],'wind':[]}); total=0; missing_t=0;missing_w=0; first=None;last=None
    for (cell,*_) in sh.iter_rows(values_only=True):
        if not isinstance(cell,str) or not cell[:4].isdigit() or 'T' not in cell:continue
        vals=next(csv.reader([cell]));
        if len(vals)<3:continue
        try:utc=dt.datetime.fromisoformat(vals[0]).replace(tzinfo=dt.timezone.utc)
        except ValueError:continue
        local=utc.astimezone(ZoneInfo('Europe/Bucharest')); day=local.date()
        t=number(vals[1]);w=number(vals[2]);total+=1
        first=min(first,day) if first else day;last=max(last,day) if last else day
        if t is None:missing_t+=1
        else:d[day]['temp'].append(t)
        if w is None:missing_w+=1
        else:d[day]['wind'].append(w)
    return d,{'total_hourly_records':total,'temperature_missing':missing_t,'wind_missing':missing_w,'first_local_date':str(first),'last_local_date':str(last)}

def main():
    p=argparse.ArgumentParser();p.add_argument('--rnmca',required=True);p.add_argument('--meteo',required=True);p.add_argument('--out',default='.');a=p.parse_args()
    out=Path(a.out);out.mkdir(parents=True,exist_ok=True)
    air,ra=read_rnmca(a.rnmca);met,rm=read_meteo(a.meteo)
    # Include RNMCA days even when every hourly PM2.5 reading is missing
    # Build calendar range from the RNMCA report itself, not only valid values
    start=dt.date.fromisoformat(ra['first_date']);end=dt.date.fromisoformat(ra['last_date'])
    rows=[]; cur=start
    while cur<=end:
        pm=air.get(cur,[]);wx=met.get(cur,{'temp':[],'wind':[]});nt=len(wx['temp']);nw=len(wx['wind']);np=len(pm)
        rows.append({'date':cur.isoformat(),'station_id':'B-12','pm25_mean_ug_m3':round(statistics.mean(pm),3) if pm else '',
                     'temperature_mean_c':round(statistics.mean(wx['temp']),3) if nt else '',
                     'wind_speed_mean_kmh':round(statistics.mean(wx['wind']),3) if nw else '',
                     'pm25_valid_hours':np,'temperature_valid_hours':nt,'wind_valid_hours':nw,
                     'all_variables_available':int(bool(np and nt and nw)),
                     'complete_pm25_24h':int(np==24)})
        cur+=dt.timedelta(days=1)
    fields=list(rows[0]);
    with (out/'B12_date_aliniate.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
    full=[r for r in rows if r['all_variables_available']]
    with (out/'B12_date_valide.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(full)
    report={'method':'RNMCA 01–24 assigned to printed local calendar day; Open-Meteo GMT/UTC converted to Europe/Bucharest including DST; daily means over available hourly values',
            'rnmca':ra,'open_meteo':rm,'calendar_days_rnmca':len(rows),'days_with_all_three_variables':len(full),
            'days_without_pm25':sum(not r['pm25_valid_hours'] for r in rows),
            'days_without_temperature':sum(not r['temperature_valid_hours'] for r in rows),
            'days_without_wind':sum(not r['wind_valid_hours'] for r in rows),
            'days_with_24_valid_pm25_hours':sum(r['complete_pm25_24h'] for r in rows),
            'days_with_fewer_than_24_valid_pm25_hours':sum(r['pm25_valid_hours']<24 for r in rows),
            'days_with_all_variables_and_24_pm25_hours':sum(r['all_variables_available'] and r['complete_pm25_24h'] for r in rows),
            'notes':['RNMCA 01–24 hour-label interpretation should be checked against portal metadata.',
                     'Open-Meteo data use GMT/UTC; local conversion causes partial boundary days and 23/25-hour DST days.',
                     'Means are calculated for any day with at least one valid hourly observation; use valid-hour columns to impose a stricter completeness threshold.',
                     'Meteorological coordinates represent a model location, not necessarily the exact B-12 station.']}
    (out/'coverage_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
